from openpyxl import Workbook

from openpyxl.styles import (
    Font,
    Alignment
)


HEADERS = [
    "Group",
    "Nội dung",
    "Hình ảnh",
    "Thời gian",
    "Trạng thái",
    "Link bài đăng"
]


def export_results(
    rows,
    path
):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Marketing Results"

    # HEADER

    for column, header in enumerate(
        HEADERS,
        1
    ):

        cell = sheet.cell(
            row=1,
            column=column,
            value=header
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # DATA

    for row_index, data in enumerate(
        rows,
        2
    ):

        for column, header in enumerate(
            HEADERS,
            1
        ):

            cell = sheet.cell(
                row=row_index,
                column=column,
                value=data.get(
                    header,
                    ""
                )
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # COLUMN WIDTH

    widths = [
        30,
        60,
        50,
        25,
        25,
        70
    ]

    for index, width in enumerate(
        widths,
        1
    ):

        sheet.column_dimensions[
            chr(64 + index)
        ].width = width

    sheet.freeze_panes = "A2"

    workbook.save(
        path
    )